from io import BytesIO

from docx import Document
from PIL import Image

from app.docx_export import build_docx, doc_layout, pre_generate_warnings
from helpers import new_session, png_bytes


def test_layout_skips_skip_without_images(tmp_path, xlsx):
    s = new_session(tmp_path, xlsx)  # TC02 = Skip, no images
    layout = doc_layout(s)
    assert [m["module"] for m in layout] == ["Login", "Transfer"]
    assert [c["tc_no"] for c in layout[0]["cases"]] == ["TC01"]
    s.set_status("TC03", "Skip")
    assert [m["module"] for m in doc_layout(s)] == ["Login"]


def test_layout_keeps_skip_with_images_and_widths(tmp_path, xlsx):
    s = new_session(tmp_path, xlsx)
    s.set_active("TC02")
    for color in [(1, 1, 1), (2, 2, 2), (3, 3, 3)]:
        s.add_image(png_bytes(color=color))
    case = doc_layout(s)[0]["cases"][1]
    assert case["tc_no"] == "TC02" and case["width_cm"] == 5.1
    assert doc_layout(s)[0]["cases"][0]["width_cm"] == 7.6


def test_warnings(tmp_path, xlsx):
    s = new_session(tmp_path, xlsx)
    s.set_status("TC02", "Todo")
    w = pre_generate_warnings(s)
    assert w["no_image"] == [
        "Login · TC001 - User input username valid",
        "Login · TC002 - User input username invalid",
        "Transfer · TC001 - User Transfer Bank A to Bank A",
    ]
    assert w["todo"] == ["Login · TC002 - User input username invalid"]


def test_build_docx_structure(tmp_path, xlsx):
    s = new_session(tmp_path, xlsx)
    s.set_active("TC01")
    s.add_image(png_bytes(1600, 3000, (10, 10, 10)))
    s.add_image(png_bytes(1600, 3000, (20, 20, 20), fmt="JPEG"), ".jpg")
    s.set_active("TC03")
    s.add_image(png_bytes(color=(30, 30, 30)))
    s.add_image(png_bytes(color=(0, 0, 255, 128), mode="RGBA"))
    s.add_image(png_bytes(color=(50, 50, 50)))

    path = build_docx(s)
    assert path.name == "Regression _ Regresi - Android 1.2.4.docx"

    doc = Document(path)
    heads = [(p.style.name, p.text) for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert heads == [
        ("Heading 1", "Login"),
        ("Heading 2", "TC001 - User input username valid"),
        ("Heading 1", "Transfer"),
        ("Heading 2", "TC001 - User Transfer Bank A to Bank A"),
    ]
    per_paragraph = [len(p._p.xpath(".//pic:pic")) for p in doc.paragraphs if p._p.xpath(".//pic:pic")]
    assert per_paragraph == [2, 3]
    assert [round(sh.width.cm, 1) for sh in doc.inline_shapes] == [7.6, 7.6, 5.1, 5.1, 5.1]
    assert doc.element.body.xml.count('w:type="page"') == 1
    for part in doc.part.package.image_parts:
        assert Image.open(BytesIO(part.blob)).width <= 800


def test_missing_image_file_is_skipped(tmp_path, xlsx):
    s = new_session(tmp_path, xlsx)
    s.set_active("TC01")
    name = s.add_image(png_bytes())
    (s.shots / name).unlink()
    assert build_docx(s).exists()


def test_images_of_a_case_fit_on_one_line(tmp_path, xlsx):
    # Word wraps a picture that does not fit the text width onto the next line (and with phone
    # screenshots, the next page), while Preview Doc shows them side by side.
    s = new_session(tmp_path, xlsx)
    for tc, count in (("TC01", 2), ("TC03", 3)):
        s.set_active(tc)
        for i in range(count):
            s.add_image(png_bytes(width=100, height=222, color=(i, i, i)))
    doc = Document(str(build_docx(s)))
    section = doc.sections[0]
    text_width = section.page_width - section.left_margin - section.right_margin
    space = 0.2 * 360000  # generous width of the space between pictures, in EMU
    rows = [p for p in doc.paragraphs if p._p.xpath(".//pic:pic")]
    assert len(rows) == 2
    for p in rows:
        widths = [int(e) for e in p._p.xpath(".//wp:extent/@cx")]
        assert sum(widths) + space * (len(widths) - 1) <= text_width


def test_sit_numbers_through_the_whole_document(tmp_path, xlsx):
    regression = new_session(tmp_path / "r", xlsx)
    sit = new_session(tmp_path / "s", xlsx, doc_type="sit")
    assert regression.labels() == {"TC01": "TC001", "TC02": "TC002", "TC03": "TC001"}
    assert sit.labels() == {"TC01": "TC001", "TC02": "TC002", "TC03": "TC003"}
    assert sit.headings()["TC03"] == "TC003 - User Transfer Bank A to Bank A"


def test_sit_docx_one_table_per_case_with_pictures_in_the_cell(tmp_path, xlsx):
    from docx.oxml.ns import qn
    from docx.shared import Twips

    s = new_session(tmp_path, xlsx, doc_type="sit")
    s.set_active("TC01")
    s.add_image(png_bytes(color=(1, 1, 1)))
    s.add_image(png_bytes(color=(2, 2, 2)))
    path = build_docx(s)
    assert path.name == "SIT _ Regresi - Android 1.2.4.docx"
    doc = Document(path)
    section = doc.sections[0]
    assert (section.page_width, section.page_height) == (Twips(12240), Twips(15840))
    assert [h.text for h in doc.paragraphs if h.style.name.startswith("Heading")] == []
    assert len(doc.tables) == 2  # TC02 is Skip without screenshots
    first, second = doc.tables
    assert [r.cells[0].text for r in first.rows] == [
        "No. TC", "Scenario", "Category", "Steps", "Expected Result", "Screen Capture"]
    assert [r.cells[1].text for r in first.rows][:5] == [
        "TC001", "User input username valid", "Happy Flow", "1. Login", "Home page"]
    assert all(r.cells[0].paragraphs[0].runs[0].bold for r in first.rows)
    tbl_w = first._tbl.tblPr.find(qn("w:tblW"))
    assert (tbl_w.get(qn("w:w")), tbl_w.get(qn("w:type"))) == ("10512", "dxa")
    assert [c.get(qn("w:w")) for c in first._tbl.tblGrid.findall(qn("w:gridCol"))] == ["1350", "9162"]
    assert len(first.rows[-1].cells[1]._element.xpath(".//pic:pic")) == 2
    assert second.rows[0].cells[1].text == "TC003"
    assert second.rows[1].cells[1].text == "User Transfer Bank A to Bank A"
    assert len(second.rows[-1].cells[1]._element.xpath(".//pic:pic")) == 0
    breaks = [p for p in doc.paragraphs if p._element.xpath('.//w:br[@w:type="page"]')]
    assert len(breaks) == 1  # between the two tables, none after the last


def test_sit_steps_and_expected_one_paragraph_per_line(tmp_path, xlsx):
    s = new_session(tmp_path, xlsx, doc_type="sit")
    case = s.case("TC01")
    case["steps"] = "1. User is on the Homepage\n\n2. User tap notification bell icon"
    case["expected"] = "- Navigated to the notification list\n- Bell is highlighted"
    doc = Document(build_docx(s))
    steps, expected = doc.tables[0].rows[3].cells[1], doc.tables[0].rows[4].cells[1]
    assert [p.text for p in steps.paragraphs] == ["1. User is on the Homepage", "2. User tap notification bell icon"]
    assert [p.text for p in expected.paragraphs] == ["- Navigated to the notification list", "- Bell is highlighted"]


def test_unknown_doc_type_is_rejected(tmp_path, xlsx):
    import pytest

    from app.session import SessionError
    with pytest.raises(SessionError):
        new_session(tmp_path, xlsx, doc_type="uat")
