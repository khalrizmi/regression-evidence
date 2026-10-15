"""Read test cases from one sheet of an Excel workbook."""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

REQUIRED = {"no tc": "tc_no", "function": "module", "scenario": "title"}
OPTIONAL = {
    "type": "type",
    "expected result": "expected",
    "step": "steps",
    "steps": "steps",
    "remark": "remark",
    "remarks": "remark",
}
PLATFORMS = ("android", "ios")
NO_MODULE = "(Tanpa Function)"


class SheetError(ValueError):
    pass


@dataclass
class TestCase:
    tc_no: str
    row: int
    module: str
    title: str
    type: str = ""
    expected: str = ""
    steps: str = ""
    remark: str = ""
    results: dict[str, str] = field(default_factory=dict)


@dataclass
class SheetData:
    sheet: str
    cases: list[TestCase]
    result_columns: dict[str, list[str]]


def normalize_header(value) -> str:
    text = str(value or "").replace(".", "")
    return re.sub(r"\s+", " ", text).strip().lower()


def cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _label(value) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def list_sheets(path: Path) -> list[str]:
    wb = load_workbook(path, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def _find_header(rows: list[list]) -> int:
    for idx, row in enumerate(rows[:10]):
        if "no tc" in [normalize_header(v) for v in row]:
            return idx
    raise SheetError("Header 'No TC' tidak ditemukan di 10 baris pertama")


def load_sheet(path: Path, sheet: str) -> SheetData:
    wb = load_workbook(path, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            raise SheetError(f"Sheet tidak ditemukan: {sheet}")
        rows = [list(r) for r in wb[sheet].iter_rows(values_only=True)]
    finally:
        wb.close()

    header_idx = _find_header(rows)
    header = rows[header_idx]
    names = [normalize_header(v) for v in header]

    cols: dict[str, int] = {}
    for idx, name in enumerate(names):
        field_name = REQUIRED.get(name) or OPTIONAL.get(name)
        if field_name and field_name not in cols:
            cols[field_name] = idx
    missing = [h for h, f in REQUIRED.items() if f not in cols]
    if missing:
        found = ", ".join(_label(h) for h in header if h)
        raise SheetError(f"Kolom wajib tidak ada: {', '.join(missing)}. Header terbaca: {found}")

    result_columns: dict[str, list[str]] = {p: [] for p in PLATFORMS}
    result_idx: dict[str, int] = {}
    for idx, name in enumerate(names):
        for platform in PLATFORMS:
            if name.startswith(f"result {platform}"):
                label = _label(header[idx])
                result_columns[platform].append(label)
                result_idx[label] = idx

    def get(row: list, idx: int | None) -> str:
        return cell_text(row[idx]) if idx is not None and idx < len(row) else ""

    cases: list[TestCase] = []
    seen: set[str] = set()
    module = ""
    for offset, row in enumerate(rows[header_idx + 1:]):
        excel_row = header_idx + 2 + offset
        tc_no = get(row, cols["tc_no"])
        title = _label(get(row, cols["title"]))
        module = get(row, cols["module"]) or module
        if not tc_no and not title:
            continue
        if not tc_no:
            raise SheetError(f"Baris {excel_row}: No TC kosong")
        if tc_no in seen:
            raise SheetError(f"No TC duplikat: {tc_no} (baris {excel_row})")
        seen.add(tc_no)
        cases.append(TestCase(
            tc_no=tc_no,
            row=excel_row,
            module=module or NO_MODULE,
            title=title,
            type=get(row, cols.get("type")),
            expected=get(row, cols.get("expected")),
            steps=get(row, cols.get("steps")),
            remark=get(row, cols.get("remark")),
            results={label: get(row, idx) for label, idx in result_idx.items()},
        ))
    return SheetData(sheet=sheet, cases=cases, result_columns=result_columns)
