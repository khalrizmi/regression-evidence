"""Evidence document: layout shared with Preview Doc, and the .docx writer."""
from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, Twips
from PIL import Image

MAX_WIDTH_PX = 800
PLATFORM_LABEL = {"android": "Android", "ios": "iOS"}
UNSAFE_FILENAME = re.compile(r'[\\/:*?"<>|]+')


MARGIN_CM = 2.5  # same as the page padding of Preview Doc (app/static/style.css .page)

# SIT layout, measured from the team's SIT evidence document: US Letter, one 2-column table per TC.
SIT_PAGE = (Twips(12240), Twips(15840))
SIT_MARGINS = (Twips(864), Twips(792))  # left/right, top/bottom
SIT_COLUMNS = (Twips(1350), Twips(9162))  # label, value: together the full text width
SIT_ROWS = ("No. TC", "Scenario", "Category", "Steps", "Expected Result", "Screen Capture")
OUTPUT_PREFIX = {"regression": "Regression", "sit": "SIT"}


def image_width_cm(count: int, doc_type: str = "regression") -> float:
    if doc_type == "sit":
        return 7.4 if count <= 2 else 5.0  # fits the 16.2 cm value column
    return 7.6 if count <= 2 else 5.1


def doc_layout(session) -> list[dict]:
    headings = session.headings()
    labels = session.labels()
    modules: dict[str, list[dict]] = {}
    for c in session.active_cases():
        if c["status"] == "Skip" and not c["images"]:
            continue
        modules.setdefault(c["module"], []).append({
            "tc_no": c["tc_no"],
            "label": labels[c["tc_no"]],
            "heading": headings[c["tc_no"]],
            "title": c["title"],
            "type": c["type"],
            "steps": c["steps"],
            "expected": c["expected"],
            "status": c["status"],
            "images": list(c["images"]),
            "width_cm": image_width_cm(len(c["images"]), session.doc_type),
        })
    return [{"module": m, "cases": cases} for m, cases in modules.items()]


def pre_generate_warnings(session) -> dict:
    headings = session.headings()
    no_image, todo = [], []
    for c in session.active_cases():
        label = f"{c['module']} · {headings[c['tc_no']]}"
        if c["status"] != "Skip" and not c["images"]:
            no_image.append(label)
        if c["status"] == "Todo":
            todo.append(label)
    return {"no_image": no_image, "todo": todo}


def shrink(path: Path) -> BytesIO:
    with Image.open(path) as img:
        img = img.convert("RGB")
        if img.width > MAX_WIDTH_PX:
            height = round(img.height * MAX_WIDTH_PX / img.width)
            img = img.resize((MAX_WIDTH_PX, height), Image.Resampling.LANCZOS)
        buf = BytesIO()
        img.save(buf, "JPEG", quality=85)
    buf.seek(0)
    return buf


def output_path(session) -> Path:
    d = session.data
    sheet = UNSAFE_FILENAME.sub("-", d["sheet"]).strip()
    build = UNSAFE_FILENAME.sub("-", d["build"]).strip()
    platform = PLATFORM_LABEL.get(d["platform"], d["platform"])
    prefix = OUTPUT_PREFIX[session.doc_type]
    return session.root / "out" / f"{prefix} _ {sheet} - {platform} {build}.docx"


def build_docx(session) -> Path:
    doc = build_sit(session) if session.doc_type == "sit" else build_regression(session)
    path = output_path(session)
    path.parent.mkdir(exist_ok=True)
    doc.save(path)
    return path


def build_regression(session) -> Document:
    doc = Document()
    # python-docx's default 3.17 cm side margins leave 15.24 cm: too narrow for 2 x 7.6 cm or 3 x 5.1 cm,
    # so Word would wrap the last picture onto the next page instead of beside the others.
    for section in doc.sections:
        section.left_margin = section.right_margin = Cm(MARGIN_CM)
        section.top_margin = section.bottom_margin = Cm(MARGIN_CM)
    layout = doc_layout(session)
    total = sum(len(m["cases"]) for m in layout)
    written = 0
    for module in layout:
        doc.add_heading(module["module"], level=1)
        for case in module["cases"]:
            doc.add_heading(case["heading"], level=2)
            files = [session.shots / n for n in case["images"] if (session.shots / n).exists()]
            if files:
                run = doc.add_paragraph().add_run()
                width = Cm(image_width_cm(len(files)))
                for i, f in enumerate(files):
                    if i:
                        run.add_text(" ")
                    run.add_picture(shrink(f), width=width)
            written += 1
            if written < total:
                doc.add_page_break()
    return doc


def _borders(table) -> None:
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        for key, value in (("val", "single"), ("sz", "6"), ("space", "0"), ("color", "000000")):
            el.set(qn(f"w:{key}"), value)
        borders.append(el)
    table._tbl.tblPr.append(borders)


def _fixed_width(table) -> None:
    """Word honours these; without them it sizes the columns to their content."""
    tbl_w = table._tbl.tblPr.find(qn("w:tblW"))
    tbl_w.set(qn("w:w"), str(sum(SIT_COLUMNS) // Twips(1)))
    tbl_w.set(qn("w:type"), "dxa")
    for col, width in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), SIT_COLUMNS):
        col.set(qn("w:w"), str(width // Twips(1)))


def _fill(cell, lines: list[str], bold: bool = False, space_after: float = 1) -> None:
    """One paragraph per line, like the SIT template (numbered steps, '- ' expected results)."""
    for i, line in enumerate(lines or [""]):
        para = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        para.paragraph_format.space_after = Pt(space_after)
        para.paragraph_format.line_spacing = 1.15
        para.add_run(line).bold = bold


def _lines(text: str) -> list[str]:
    return [line.rstrip() for line in (text or "").splitlines() if line.strip()]


def build_sit(session) -> Document:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(9)
    for section in doc.sections:
        section.page_width, section.page_height = SIT_PAGE
        section.left_margin = section.right_margin = SIT_MARGINS[0]
        section.top_margin = section.bottom_margin = SIT_MARGINS[1]
    cases = [case for module in doc_layout(session) for case in module["cases"]]
    for n, case in enumerate(cases, 1):
        table = doc.add_table(rows=len(SIT_ROWS), cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        _fixed_width(table)
        _borders(table)
        values = [[case["label"]], [case["title"]], [case["type"]], _lines(case["steps"]),
                  _lines(case["expected"]), None]
        for row, label, value in zip(table.rows, SIT_ROWS, values):
            for cell, width in zip(row.cells, SIT_COLUMNS):
                cell.width = width
            _fill(row.cells[0], [label], bold=True, space_after=10)  # the template's row padding
            if value is not None:
                _fill(row.cells[1], value)
        files = [session.shots / name for name in case["images"] if (session.shots / name).exists()]
        if files:
            para = table.rows[-1].cells[1].paragraphs[0]
            run = para.add_run()
            width = Cm(image_width_cm(len(files), "sit"))
            for i, f in enumerate(files):
                if i:
                    run.add_text(" ")
                run.add_picture(shrink(f), width=width)
        if n < len(cases):
            doc.add_page_break()
    return doc
